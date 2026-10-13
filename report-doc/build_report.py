"""สร้างเล่มรายงาน SA วิชา 01418321 (.docx) รูปแบบเดียวกับรายงาน 01418344
ใช้:  python3 report-doc/build_report.py
แล้วอัปเดตสารบัญและสร้าง PDF ด้วย:  ./report-doc/export_pdf.sh

เนื้อหาอยู่ใน content/ (ch1.py, ch2.py, usecases.json, tables.json)
รูปอยู่ในโฟลเดอร์ของแต่ละบท ส่วน SQL และ diagram ของ 2.14-2.21, 2.25 สร้างจาก 04-chapter2-diagrams-database
"""
import io
import json
import os
import re
import struct
import sys

from docx.text.paragraph import Paragraph
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FONT = "TH Sarabun New"
CODE_FONT = "Menlo"
NAME = "รายงาน-SA-Colorado"
# PRINT=1 : แทรกจุดตัดคำที่มองไม่เห็นระหว่างคำไทย เพื่อให้ Word (Mac) ตัดบรรทัดถูกตอน export PDF
# PRINT=0 : ไม่แทรก สำหรับนำเข้า Google Docs
PRINT = os.environ.get("PRINT", "1") == "1"
OUT = os.path.join(HERE, "output", NAME + ".docx") if PRINT else os.path.join(HERE, "output", "สำหรับ-Google-Docs", NAME + ".docx")
TITLE = "ระบบจัดการหลังบ้านอัตโนมัติ Colorado Co., Ltd."
SUBTITLE = "เฟส 1: Rakuten Ichiba, Rakuten RMS และ RSL"
MEMBERS = [("6710450911", "นายธนภัทร พูลโภคา"), ("6710451488", "นายอชิระ ภู่รุ่งเรือง"),
           ("6710451054", "นางสาวปิยากร ผลพานิช")]
SECTION_NO = "หมู่ 200"

CONTENT = os.path.join(HERE, "content")
BP_IMG = lambda n: os.path.join(ROOT, "02-chapter2-business-process", "img", "final", n)
DIA_IMG = lambda n: os.path.join(ROOT, "03-chapter2-use-case-descriptions", "diagrams", "png", n)
DB_IMG = lambda n: os.path.join(ROOT, "04-chapter2-diagrams-database", "diagrams", "png", n)
DB_DATA = lambda n: os.path.join(ROOT, "04-chapter2-diagrams-database", "sql", n)
PLT_IMG = lambda n: os.path.join(ROOT, "05-chapter2-platform-testing", "img", n)
LOGO = os.path.join(HERE, "assets", "ku_logo.png")
RUBRIC = os.path.join(HERE, "รายละเอียดของคะแนนโครงงาน.xlsx")

if PRINT:
    from pythainlp.tokenize import word_tokenize
    _TH = re.compile(r"[฀-๿]")

    def brk(text):
        out = []
        for tok in word_tokenize(text, engine="newmm", keep_whitespace=True):
            if out and _TH.search(tok) and _TH.search(out[-1]) and not out[-1].isspace():
                out.append("​")
            out.append(tok)
        return "".join(out)
else:
    brk = lambda t: t

doc = Document()


# ------------------------------------------------------------------ styles
def set_font(style, size, bold=False):
    f = style.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.cs_bold = bold; f.color.rgb = None
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    sz = rpr.find(qn("w:szCs"))
    if sz is None: sz = OxmlElement("w:szCs"); rpr.append(sz)
    sz.set(qn("w:val"), str(size * 2))
    c = rpr.find(qn("w:color"))
    if c is not None:
        c.set(qn("w:val"), "000000"); [c.attrib.pop(qn(k), None) for k in ("w:themeColor", "w:themeShade", "w:themeTint")]


S = doc.styles
set_font(S["Normal"], 16)
S["Normal"].paragraph_format.space_after = Pt(0); S["Normal"].paragraph_format.line_spacing = 1.0
for name, size, before, after in (("Heading 1", 20, 18, 10), ("Heading 2", 18, 18, 8), ("Heading 3", 16, 14, 6), ("Title", 22, 0, 0)):
    set_font(S[name], size, True)
    pf = S[name].paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.keep_with_next = True
set_font(S["Caption"], 16, True); S["Caption"].font.italic = False
for n, sid in (("toc 1", "TOC1"), ("toc 2", "TOC2"), ("toc 3", "TOC3"), ("table of figures", "TableofFigures")):
    st = S.add_style(n, 1); st.base_style = S["Normal"]; st.element.set(qn("w:styleId"), sid)
    st.element.find(qn("w:name")).set(qn("w:val"), n)
    if n == "toc 2": st.paragraph_format.left_indent = Inches(0.3)
    if n == "toc 3": st.paragraph_format.left_indent = Inches(0.6)
    set_font(st, 16); st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 0.92


def set_thai_lang(rpr):
    lg = rpr.find(qn("w:lang"))
    if lg is None: lg = OxmlElement("w:lang"); rpr.append(lg)
    lg.set(qn("w:val"), "en-US"); lg.set(qn("w:eastAsia"), "en-US"); lg.set(qn("w:bidi"), "th-TH")


set_thai_lang(doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")))
for st in S:
    if st.type in (1, 2) and st.element.find(qn("w:rPr")) is not None: set_thai_lang(st.element.find(qn("w:rPr")))
_set = doc.settings.element
tfl = _set.find(qn("w:themeFontLang"))
if tfl is None: tfl = OxmlElement("w:themeFontLang"); _set.append(tfl)
tfl.set(qn("w:val"), "en-US"); tfl.set(qn("w:bidi"), "th-TH"); tfl.set(qn("w:eastAsia"), "en-US")
compat = _set.find(qn("w:compat"))
if compat is not None:
    fe = compat.find(qn("w:useFELayout"))
    if fe is not None: compat.remove(fe)
    for cs in compat.findall(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode": cs.set(qn("w:val"), "15")

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, side, Inches(1))


# ------------------------------------------------------------------ helpers
def field(par, instr, placeholder=""):
    r = par.add_run(); b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    r = par.add_run(); t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr; r._r.append(t)
    r = par.add_run(); s_ = OxmlElement("w:fldChar"); s_.set(qn("w:fldCharType"), "separate"); r._r.append(s_)
    par.add_run(placeholder)
    r = par.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)


def run(par, text, bold=False, size=None, code=False):
    r = par.add_run(text if code else brk(text))
    if bold: r.font.bold = True; r.font.cs_bold = True
    if code:
        r.font.name = CODE_FONT
        r._r.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:cs"), FONT)
        size = (size or 16) * 0.7
    if size:
        r.font.size = Pt(size); sz = OxmlElement("w:szCs"); sz.set(qn("w:val"), str(int(size * 2))); r._r.get_or_add_rPr().append(sz)
    return r


_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def rich(par, text, size=None, bold=False):
    """แปลง markdown แบบ inline: **ตัวหนา**, `โค้ด`, *ตัวเอียง* (ตัด * ทิ้ง)"""
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    for part in _INLINE.split(text):
        if not part: continue
        if part.startswith("**"): rich(par, part[2:-2], size, True)
        elif part.startswith("`"): run(par, part[1:-1], bold, size, code=True)
        else: run(par, part, bold, size)


BODY_LINE = 1.15


def P(text, indent=True, align="j"):
    p = doc.add_paragraph()
    p.alignment = {"j": WD_ALIGN_PARAGRAPH.THAI_JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if indent: p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = BODY_LINE
    rich(p, text); return p


def H(level, text): return doc.add_heading(brk(text), level)


def LI(text, marker="•", level=0):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.left_indent = Inches(0.85 + 0.4 * level); pf.first_line_indent = Inches(-0.35)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT; pf.tab_stops.add_tab_stop(Inches(0.85 + 0.4 * level))
    pf.space_after = Pt(4); pf.line_spacing = BODY_LINE
    rich(p, f"{marker}\t{text}"); return p


def CODE(text):
    for line in text.splitlines():
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5)
        run(p, line, size=16, code=True)
    doc.add_paragraph()


def png_size(path):
    with open(path, "rb") as f: head = f.read(24)
    return struct.unpack(">II", head[16:24])


def shrink(path, max_px=3200):
    """ย่อรูปใหญ่ก่อนใส่ ไม่ให้ไฟล์ docx ใหญ่เกินไป"""
    w, h = png_size(path)
    if max(w, h) <= max_px: return path
    from PIL import Image
    im = Image.open(path); im.thumbnail((max_px, max_px))
    buf = io.BytesIO(); im.save(buf, "PNG", optimize=True); buf.seek(0)
    return buf


state = {"landscape": False}


def set_orientation(landscape):
    if state["landscape"] == landscape: return
    # ย่อ paragraph ว่างท้าย section (บรรทัดเว้นหลังตาราง และตัวที่ถือ section break) ให้เหลือ 1pt
    # ไม่งั้นถ้าหน้าก่อนเต็มพอดีจะได้หน้าว่าง
    tail = doc.element.body.xpath("./w:p[last()]")
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    for el in [e for e in tail if not e.xpath(".//w:t|.//w:drawing")] + doc.element.body.xpath("./w:p[last()]"):
        pf = Paragraph(el, doc._body).paragraph_format
        pf.space_before = pf.space_after = Pt(0); pf.line_spacing = Pt(1)
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    state["landscape"] = landscape; _pb["on"] = False


fig_no = {"n": {}}


def best_landscape(path):
    """รูปนี้วางแนวนอนแล้วใหญ่กว่าแนวตั้งหรือไม่"""
    w, h = png_size(path)
    return min(9.0 / w, 5.4 / h) > 1.1 * min(6.27 / w, 7.8 / h)


def natural_w(path):
    """ความกว้างสูงสุดของรูปจาก mermaid (scale 2, ตัวอักษร 16px) ไม่ให้ตัวอักษรใหญ่เกิน 10pt"""
    return png_size(path)[0] / 2 * 10 / 16 / 72


def AFIG(path, caption, chapter, **kw):
    """ใส่รูปเต็มหน้า เลือกแนวกระดาษตามสัดส่วนรูป"""
    set_orientation(best_landscape(path))
    return FIG(path, caption, chapter, **{"max_h": 7.8, "land_h": 5.6, **kw})


def FIG(path, caption, chapter, max_h=8.2, width=6.27, as_table=False, land_h=5.4, max_w=None):
    """as_table=True : รูปที่เป็นตาราง (CRUD) ใส่คำบรรยายแบบตารางไว้เหนือรูป"""
    if not os.path.exists(path):
        P(f"[ไม่พบรูป {os.path.basename(path)}]", indent=False, align="c"); return None
    if as_table: cap = CAPTION("ตารางที่", caption, chapter, tab_no)
    w, h = png_size(path)
    mw, mh = (9.0, land_h) if state["landscape"] else (width, max_h)
    scale = min(mw / w, mh / h, (max_w or mw) / w)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = not as_table
    p.paragraph_format.space_before = Pt(4 if as_table else 10)
    p.add_run().add_picture(shrink(path), width=Inches(w * scale))
    if as_table: doc.add_paragraph(); return cap
    return CAPTION("รูปที่", caption, chapter, fig_no)


def CAPTION(kind, caption, chapter, counter):
    k = counter["n"].get(chapter, 0) + 1; counter["n"][chapter] = k
    c = doc.add_paragraph(style="Caption"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if kind == "รูปที่": c.paragraph_format.space_before = Pt(6); c.paragraph_format.space_after = Pt(14)
    else: c.paragraph_format.space_before = Pt(12); c.paragraph_format.space_after = Pt(4); c.paragraph_format.keep_with_next = True
    run(c, f"{kind} {chapter}-{k} {caption}")
    return f"{chapter}-{k}"


tab_no = {"n": {}}


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_); tcPr.append(s)


def cell_text(cell, text, size, bold=False):
    """ข้อความในเซลล์ แยกบรรทัดตาม <br>"""
    lines = re.split(r"<br\s*/?>", str(text))
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rich(par, line.strip(), size, bold)


def cell_margins(tb, v, h):
    """ระยะห่างข้อความกับเส้นตาราง หน่วย twip"""
    tblPr = tb._tbl.tblPr; mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", v), ("left", h), ("bottom", v), ("right", h)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); mar.append(e)
    look = tblPr.find(qn("w:tblLook"))
    if look is not None: look.addprevious(mar)
    else: tblPr.append(mar)


def TABLE(rows, widths, caption, chapter, size=14, header_fill="1868DB", pad=50):
    CAPTION("ตารางที่", caption, chapter, tab_no)
    tb = doc.add_table(rows=len(rows), cols=len(widths)); tb.style = "Table Grid"
    cell_margins(tb, pad, 110)
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    for c, wd in enumerate(widths): tb.columns[c].width = Inches(wd)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tb.cell(r, c); cell.width = Inches(widths[c])
            cell_text(cell, val, size, bold=(r == 0))
            if r == 0:
                shade(cell, header_fill)
                for rr in cell.paragraphs[0].runs: rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if len(rows) > 1:  # หัวตารางซ้ำเมื่อขึ้นหน้าใหม่
        trPr = tb.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    doc.add_paragraph()


_pb = {"on": False}
def page_break(): _pb["on"] = True
_add_par, _add_head = doc.add_paragraph, doc.add_heading
def _wrap(fn):
    def inner(*a, **k):
        p = fn(*a, **k)
        if _pb["on"]: p.paragraph_format.page_break_before = True; _pb["on"] = False
        return p
    return inner
doc.add_paragraph, doc.add_heading = _wrap(_add_par), _wrap(_add_head)


# ------------------------------------------------------------------ เนื้อหา
sys.path.insert(0, HERE)
from content import ch1, ch2  # noqa: E402

load = lambda p: json.load(open(p, encoding="utf-8"))
UCS = load(os.path.join(CONTENT, "usecases.json"))
DICT = load(DB_DATA("dictionary.json"))
QUERIES = load(DB_DATA("queries.json"))
CRUD = load(DB_DATA("crud.json"))
RESULTS = {r["id"]: r for r in load(DB_DATA("sql-results.json"))}


def blocks(items, chapter):
    n = 0
    for b in items:
        kind = b[0]
        if kind == "h2": H(2, b[1])
        elif kind == "h3": H(3, b[1])
        elif kind == "p": P(b[1])
        elif kind == "li": LI(b[1])
        elif kind == "ol":
            n += 1; LI(b[1], marker=f"{n}.")
        elif kind == "code": CODE(b[1])
        elif kind == "table": TABLE(b[3], b[2], b[1], chapter)
        if kind != "ol": n = 0 if kind in ("h2", "h3", "p") else n


# ------------------------------------------------------------------ ปก
def cover(text, size=18, bold=True, before=0, after=0):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, text, bold, size)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after); return p


p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(LOGO, width=Inches(1.35))
cover("รายงานวิชา 01418321", 24, before=12)
cover("System Analysis and Design", 20, before=2)
cover("การวิเคราะห์และการออกแบบระบบ", 20)
cover("เรื่อง", 20, False, before=30)
cover(TITLE, 24, before=2)
cover(SUBTITLE, 20)
cover("เสนอ", 20, False, before=30)
cover("อาจารย์ สมโชค เรืองอิทธินันท์", 20)
cover("จัดทำโดย", 20, False, before=30, after=2)
for sid, nm in MEMBERS:
    p = cover("", 18, False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.left_indent = Inches(1.1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(2.4)); p.paragraph_format.tab_stops.add_tab_stop(Inches(5.0))
    run(p, f"{sid}\t{nm}\t{SECTION_NO}", False, 18)
cover("รายงานนี้เป็นส่วนหนึ่งของรายวิชา 01418321 การวิเคราะห์และการออกแบบระบบ", 18, False, before=30)
cover("ภาควิชาวิทยาการคอมพิวเตอร์ คณะวิทยาศาสตร์ มหาวิทยาลัยเกษตรศาสตร์", 18, False)
cover("ภาคต้น ปีการศึกษา 2569", 18, False)
page_break()


def FRONT(text, toc=True):
    if toc: h = H(1, text)
    else:
        h = doc.add_paragraph(); run(h, text, True, 20)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.paragraph_format.space_after = Pt(12); return h


FRONT("สมาชิกในกลุ่ม", toc=False)
tb = doc.add_table(rows=2, cols=len(MEMBERS)); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, (sid, nm) in enumerate(MEMBERS):
    ph = tb.cell(0, c).paragraphs[0]; ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(6): ph.add_run().add_break()
    run(ph, "[ใส่รูป]")
    for _ in range(6): ph.add_run().add_break()
    pn = tb.cell(1, c).paragraphs[0]; pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(pn, nm, size=16); pn.add_run().add_break(); run(pn, f"{sid} {SECTION_NO}", size=16)
page_break()


def keywords(label, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    rich(p, f"**{label}** {text}")


FRONT("บทคัดย่อ")
for t in ch1.ABSTRACT_TH: P(t)
keywords("คำสำคัญ:", ch1.KEYWORDS_TH)
page_break()
FRONT("Abstract")
for t in ch1.ABSTRACT_EN: P(t)
keywords("Keywords:", ch1.KEYWORDS_EN)
page_break()

FRONT("สารบัญ", toc=False)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(p, "หน้า", True)
field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "(คลิกขวาแล้วเลือก Update Field เพื่อสร้างสารบัญ)")
page_break()
FRONT("สารบัญรูปและตาราง", toc=False)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(p, "หน้า", True)
field(doc.add_paragraph(), 'TOC \\h \\z \\t "Caption,1"', "(คลิกขวาแล้วเลือก Update Field เพื่อสร้างสารบัญรูปและตาราง)")
page_break()

# ------------------------------------------------------------------ บทที่ 1
H(1, "บทที่ 1 ที่มาและความสำคัญ")
blocks(ch1.CH1, 1)

# ------------------------------------------------------------------ บทที่ 2: 2.1-2.12
page_break()
H(1, "บทที่ 2 การวิเคราะห์และออกแบบระบบ")
H(2, "2.1 Business Process เดิม (As-Is)")
FIG(BP_IMG("2.1-business-process-as-is.png"), "Business Process เดิม (As-Is)", 2)
H(2, "2.2 คำอธิบาย Business Process เดิม")
for t in ch2.BP_ASIS_DESC: P(t)
H(2, "2.3 ปัญหาของ Business Process เดิม")
FIG(BP_IMG("2.3-business-process-pain-points.png"), "จุดที่เป็นปัญหาใน Business Process เดิม", 2)
for t in ch2.BP_PROBLEMS: P(t)
H(2, "2.4 Business Process ใหม่ (To-Be)")
FIG(BP_IMG("2.4-business-process-to-be.png"), "Business Process ใหม่ (To-Be)", 2)
H(2, "2.5 คำอธิบาย Business Process ใหม่")
P(ch2.BP_TOBE_INTRO)
TABLE(ch2.BP_TOBE_STEPS, (0.6, 1.9, 2.6, 1.17), "ขั้นตอนของ Business Process ใหม่ ผู้ทำ และปัญหาที่แก้", 2, size=13)
P(ch2.BP_TOBE_NOTE)
H(2, "2.6 Business Process ที่สัมพันธ์กับ Use Case")
FIG(BP_IMG("2.6-business-process-rel-use-case.png"), "Business Process ที่สัมพันธ์กับ Use Case", 2)
H(2, "2.7 ตารางจับคู่ระหว่าง Business Process และ Use Case Diagram")
TABLE(ch2.BP_UC_MAP, (0.6, 4.17, 1.5), "ตารางจับคู่ระหว่าง Business Process และ Use Case Diagram", 2)
for t in ch2.BP_UC_NOTES: P(t)
H(2, "2.8 Use Case Diagram ก่อนปรับปรุง")
FIG(BP_IMG("2.8-use-case-diagram-as-is.png"), "Use Case Diagram ก่อนปรับปรุง", 2)
H(2, "2.9 CRUD Table ก่อนปรับปรุง")
TABLE(ch2.CRUD_ASIS, (0.6, 2.37, 0.55, 0.55, 0.55, 0.55, 0.55, 0.55), "CRUD Table ก่อนปรับปรุง", 2, size=12)
P(ch2.CRUD_ASIS_NOTE, indent=False, align="l")
H(2, "2.10 คำอธิบาย CRUD Table ก่อนปรับปรุง")
for t in ch2.CRUD_ASIS_DESC: P(t)
H(2, "2.11 Use Case Diagram ที่ทบทวนปรับปรุงแล้ว")
FIG(BP_IMG("2.11-use-case-diagram-to-be.png"), "Use Case Diagram ที่ทบทวนปรับปรุงแล้ว", 2)
for t in ch2.UC_ACTORS: LI(t)

H(2, "2.12 CRUD Table ที่ทบทวนแล้ว")
P(ch2.CRUD_TOBE_INTRO)
for t in ch2.CRUD_TOBE_ADDED: LI(t)
set_orientation(True)
tables = CRUD["tables"]
head = ["Use Case"] + [t.replace("_", "_​") for t in tables]
rows = [head]
for uc in UCS:
    cells = CRUD["crud"].get(uc["id"], {"cells": {}})["cells"]
    rows.append([f"{uc['id']} {uc['name']}"] + [cells.get(t, "") for t in tables])
TABLE(rows, [2.6] + [round(7.0 / len(tables), 2)] * len(tables), "CRUD Table ที่ทบทวนแล้ว (สร้างจาก SQL ใน Use Case Description)", 2, size=10, pad=15)
set_orientation(False)

# ------------------------------------------------------------------ 2.13
UI = {"1A": "1A-LoginScreen", "4A": "0-DashboardScreen", "2A": "2A-OrderVerifyScreen", "1S": "1S-RslMatchScreen",
      "3A": "3A-ProductScreen", "3S": "3S-StockCheckScreen", "4S": "4S-CostCalculatorScreen",
      "5A": "5A-6A-ReorderDecisionScreen", "6A": "5A-6A-ReorderDecisionScreen", "7A": "7A-CancelOrderScreen",
      "8A": "8A-ShipmentScreen", "5S": "8A-ShipmentScreen", "9A": "9A-LabelPrintScreen", "10A": "10A-DataCleanupScreen"}

page_break()
H(2, "2.13 Use Case Description, Sequence Diagram, Collaboration Diagram และ User Interface")
n = 0
for uc in UCS:
    uid, uname = uc["id"], uc["name"]
    n += 1; H(3, f"2.13.{n} Use Case Description ที่ {uid}: {uname}")
    for label, key in (("Goal", "goal"), ("Pre-Condition", "pre"), ("Post-Condition", "post"), ("Actor", "actor")):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; rich(p, f"**{label}:** {uc[key]}")
        p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = BODY_LINE
    TABLE([("ลำดับ", "User Action", "System Response")] + uc["steps"], (0.75, 1.75, 3.77),
          f"Use Case Description ที่ {uid}: {uname}", 2, size=13)
    seq, col = DIA_IMG(f"{uid}-sequence.png"), DIA_IMG(f"{uid}-collaboration.png")
    set_orientation(best_landscape(seq))
    n += 1; H(3, f"2.13.{n} Sequence Diagram ของ Use Case Description ที่ {uid}")
    FIG(seq, f"Sequence Diagram ของ Use Case Description ที่ {uid}", 2, max_h=7.6, land_h=5.2, max_w=natural_w(seq))
    set_orientation(False)
    n += 1; H(3, f"2.13.{n} Collaboration Diagram ของ Use Case Description ที่ {uid}")
    FIG(col, f"Collaboration Diagram ของ Use Case Description ที่ {uid}", 2, max_h=6.6, max_w=natural_w(col))
    n += 1; H(3, f"2.13.{n} User Interface ของ Use Case Description ที่ {uid}")
    if uid in UI: FIG(PLT_IMG(f"ui/{UI[uid]}.png"), f"User Interface ของ Use Case Description ที่ {uid} ({UI[uid].split('-')[-1]})", 2)
    else: P("Use Case นี้ระบบทำเองโดยไม่มีหน้าจอของตัวเอง ผลการทำงานแสดงเป็นข้อความแจ้งเตือนในหน้าที่ Admin เปิดอยู่")
    page_break()

# ------------------------------------------------------------------ 2.14-2.19
CLASS_PARTS = ["เข้าสู่ระบบ", "รับและตรวจสอบคำสั่งซื้อ", "จัดการสินค้าและสต๊อก", "คำนวณต้นทุนและสั่งซื้อเติม",
               "จัดการใบปะ", "จัดส่งและแจ้งเลขติดตาม", "ลบข้อมูลเก่า", "คลาสข้อมูล (Entity)"]
set_orientation(best_landscape(DB_IMG("class-1.png")))
H(2, "2.14 Class Diagram")
for t in ch2.DB_TEXT["2.14"]: P(t)
for i, title in enumerate(CLASS_PARTS, 1):
    AFIG(DB_IMG(f"class-{i}.png"), f"Class Diagram กลุ่ม{title} ({i}/{len(CLASS_PARTS)})", 2, land_h=3.3 if i == 1 else 5.6)
set_orientation(False)

H(2, "2.15 Data Flow Diagram")
for t in ch2.DB_TEXT["2.15"]: P(t)
FIG(DB_IMG("dfd-context.png"), "Context Diagram ของระบบ", 2)
AFIG(DB_IMG("dfd-level0.png"), "Data Flow Diagram Level 0", 2)
set_orientation(False)

H(2, "2.16 State Diagram ของ Order")
for t in ch2.DB_TEXT["2.16"]: P(t)
FIG(DB_IMG("state-1.png"), "State Diagram ของ Order ก่อนจัดส่ง (1/2)", 2, max_h=6.6, max_w=natural_w(DB_IMG("state-1.png")))
FIG(DB_IMG("state-2.png"), "State Diagram ของ Order ช่วงจัดส่ง (2/2)", 2, max_h=7.8, max_w=natural_w(DB_IMG("state-2.png")))

page_break()
H(2, "2.17 Component Diagram")
for t in ch2.DB_TEXT["2.17"]: P(t)
FIG(DB_IMG("component.png"), "Component Diagram ของระบบ", 2, max_h=7.0)

page_break()
H(2, "2.18 ER Diagram")
for t in ch2.DB_TEXT["2.18"]: P(t)
FIG(DB_IMG("er-0.png"), "ER Diagram ภาพรวมของระบบ", 2, max_h=6.8)
ER_PARTS = ["คำสั่งซื้อ (ORDERS)", "ผู้ใช้ การจัดส่ง และบันทึกการลบ", "สินค้า สต๊อก ต้นทุน และการสั่งซื้อเติม"]
for i, title in enumerate(ER_PARTS, 1):
    AFIG(DB_IMG(f"er-{i}.png"), f"ER Diagram ส่วน{title} ({i}/{len(ER_PARTS)})", 2)

set_orientation(best_landscape(DB_IMG("er-query-1.png")))
H(2, "2.19 ER Diagram พร้อม Query")
for t in ch2.DB_TEXT["2.19"]: P(t)
for i, title in enumerate(ER_PARTS, 1):
    AFIG(DB_IMG(f"er-query-{i}.png"), f"ER Diagram พร้อม Query ส่วน{title} ({i}/{len(ER_PARTS)})", 2, land_h=4.0 if i == 1 else 5.6)
set_orientation(False)

# ------------------------------------------------------------------ 2.20 Query Table
H(2, "2.20 Query Table")
for t in ch2.DB_TEXT["2.20"]: P(t)
rows = [("Use Case ID", "Use Case", "Query ID", "SQL Command")]
for q in QUERIES: rows.append((q["uc"], q["uc_name"], q["id"], f"`{q['sql']}`"))
TABLE(rows, (0.7, 1.45, 0.8, 3.32), "Query Table", 2, size=12)

# ------------------------------------------------------------------ 2.21 Table Structure
page_break()
H(2, "2.21 Table Structure")
for t in ch2.DB_TEXT["2.21"]: P(t)
for i, d in enumerate(DICT, 1):
    H(3, f"2.21.{i} {d['title']}")
    rows = [("No.", "Attribute", "Description", "Data Type", "Key", "Not Null", "ช่วงข้อมูล / การตรวจสอบ")]
    for c in d["columns"]:
        rows.append((str(c["no"]), c["name"], c["desc"], c["type"], c["key"], c["notnull"], c["range"]))
    TABLE(rows, (0.35, 1.05, 1.3, 0.9, 0.65, 0.5, 1.52), f"Data Dictionary ของ{d['title']}", 2, size=11)
    fields, sample = d["sample"]["fields"], d["sample"]["rows"]
    for k in range(0, len(fields), 6):
        cols = fields[k:k + 6]
        rows = [tuple(cols)] + [tuple(r[k:k + 6]) for r in sample]
        part = f" (คอลัมน์ {k + 1}-{k + len(cols)})" if len(fields) > 6 else ""
        TABLE(rows, [round(6.27 / len(cols), 2)] * len(cols), f"ตัวอย่างข้อมูลของ{d['title']}{part}", 2, size=10)

# ------------------------------------------------------------------ 2.22-2.24
page_break()
H(2, "2.22 Software Development Platform")
FIG(PLT_IMG("2.22-software-platform.png"), "เครื่องมือที่ใช้พัฒนาระบบ (Stack of Tools)", 2)
TABLE(ch2.PLATFORM, (1.2, 2.6, 2.47), "เครื่องมือที่ใช้พัฒนาระบบ", 2)

H(2, "2.23 Site Map / UI Structure")
FIG(PLT_IMG("sitemap.png"), "โครงสร้างหน้าจอของระบบ (Site Map)", 2)
P(ch2.SITEMAP_DESC[0]); P(ch2.SITEMAP_DESC[1])
CODE(ch2.SITEMAP_TREE)
P(ch2.SITEMAP_DESC[2])
H(3, "2.23.1 User Interface Design")
P("ชื่อหน้าจอตรงกับชื่อ Screen object ใน Sequence Diagram")
TABLE([r[:2] for r in ch2.UI_SCREENS], (2.4, 3.87), "หน้าจอของระบบและ Use Case ที่เกี่ยวข้อง", 2)
for scr, ucs, img in ch2.UI_SCREENS[1:]:
    FIG(PLT_IMG(img), f"{scr} ({ucs})", 2)

H(2, "2.24 Output Design (Report)")
P("ระบบมี 2 รายงาน ทั้งคู่เป็นหน้าสำหรับพิมพ์ที่เปิดจากหน้าจอเดิม ตอนพิมพ์จะซ่อนเมนูและปุ่มออก")
H(3, "2.24.1 ใบปะสินค้า (Shipping Label)")
FIG(PLT_IMG("report/label-ORD-24178.png"), "ใบปะสินค้าของ ORD-24178", 2)
P("ใบปะสินค้าสร้างโดย Use Case 6S จัดรูปแบบใบปะสินค้า และพิมพ์ผ่าน Use Case 9A จากหน้า LabelPrintScreen ข้อมูลมาจาก Q6S.2 (ตาราง label) และตาราง orders")
TABLE(ch2.REPORT_LABEL, (1.5, 3.17, 1.6), "โครงสร้างของใบปะสินค้า", 2)
H(3, "2.24.2 รายงานต้นทุนต่อหน่วย (Unit Cost Report)")
FIG(PLT_IMG("report/unit-cost-FSH-MLR-240.png"), "รายงานต้นทุนต่อหน่วยของ FSH-MLR-240", 2)
P("รายงานต้นทุนต่อหน่วยมาจาก Use Case 4S คำนวณ Cost ในการสั่งสินค้าเติม stock เอง เปิดจากหน้า CostCalculatorScreen ใช้สูตรตาม Q4S.2 และบันทึกผลด้วย Q4S.3")
CODE(ch2.UNIT_COST_FORMULA)
TABLE(ch2.REPORT_COST, (1.6, 4.67), "โครงสร้างของรายงานต้นทุนต่อหน่วย", 2)

# ------------------------------------------------------------------ 2.25 SQL Statement
page_break()
H(2, "2.25 SQL Statement ใน Use Case Description")
for t in ch2.DB_TEXT["2.25"]: P(t)


def result_text(res):
    if not res: return ""
    lines = [", ".join(f"{f}={v}" for f, v in zip(res["fields"], row)) for row in res["rows"][:4]]
    more = f"<br/>(ทั้งหมด {res['total']} แถว)" if res["total"] > 4 else ""
    return "<br/>".join(lines) + more if lines else "ไม่พบข้อมูล (0 แถว)"


for i, uc in enumerate(UCS, 1):
    H(3, f"2.25.{i} Use Case {uc['id']}: {uc['name']}")
    rows = [("Query", "SQL Command", "ค่าที่ส่ง", "ผลการทำงาน")]
    for q in [q for q in QUERIES if q["uc"] == uc["id"]]:
        r = RESULTS[q["id"]]
        sql = f"`{q['sql']}`" + (f"<br/>ทดสอบด้วย: `{r['run_sql']}`" if r.get("run_sql") else "")
        if r.get("error"): out = f"ผิดพลาด: {r['error']}"
        elif r.get("check"): out = f"แก้ไข {r['affected']} แถว<br/>หลังรัน: " + result_text(r["check"])
        elif r.get("result"): out = result_text(r["result"])
        else: out = f"แก้ไข {r['affected']} แถว"
        rows.append((q["id"], sql, ", ".join(r["params"]) or "-", out))
    TABLE(rows, (0.65, 2.55, 1.2, 1.87), f"SQL Statement และผลการทำงานของ Use Case {uc['id']}", 2, size=10)

# ------------------------------------------------------------------ 2.26
page_break()
H(2, "2.26 การทดสอบโปรแกรมทำงานตาม Use Case Description")
for t in ch2.TEST_INTRO: P(t)
tests = ch2.TESTS[1:]
order = []
for r in tests:
    if r[0] not in order: order.append(r[0])
for k, uid in enumerate(order, 1):
    rows = [r for r in tests if r[0] == uid]
    H(3, f"2.26.{k} การทดสอบ Use Case {uid}: {rows[0][1]}")
    TABLE([("กรณีทดสอบ", "ผลที่คาดหวัง (ข้อความจาก UC)", "ผล")] + [(r[2], r[3], r[4]) for r in rows],
          (2.2, 3.37, 0.7), f"การทดสอบ Use Case {uid}: {rows[0][1]}", 2, size=14)
    seen = set()
    for r in rows:
        if r[5] in seen: continue
        seen.add(r[5])
        kind = "กรณีปกติ" if r[5].endswith("-ok.png") else "กรณีผิดพลาด"
        FIG(PLT_IMG(r[5]), f"ผลการทดสอบ Use Case {uid} {kind}", 2, max_h=4.2)
k = len(order) + 1
H(3, f"2.26.{k} การทดสอบการออกรายงาน")
TABLE([("รายงาน", "Use Case", "วิธีทดสอบ", "ผล"),
       ("ใบปะสินค้า", "6S, 9A", "จับคู่ ORD-24178 กับ RSL แล้วเปิดหน้าพิมพ์ใบปะ ข้อมูลบนใบปะต้องตรงกับ Order และเลข RSL", "ผ่าน"),
       ("รายงานต้นทุนต่อหน่วย", "4S", "คำนวณต้นทุนของ FSH-MLR-240 แล้วเปิดรายงาน ต้นทุนต่อหน่วยต้องตรงกับสูตร Q4S.2", "ผ่าน")],
      (1.5, 0.8, 3.27, 0.7), "การทดสอบการออกรายงาน", 2, size=14)
FIG(PLT_IMG("report/label-ORD-24178.png"), "ผลการออกรายงานใบปะสินค้า", 2, max_h=4.2)
FIG(PLT_IMG("report/unit-cost-FSH-MLR-240.png"), "ผลการออกรายงานต้นทุนต่อหน่วย", 2, max_h=4.2)

# ------------------------------------------------------------------ 2.27 หน้าที่รับผิดชอบ
page_break()
H(2, "2.27 หน้าที่รับผิดชอบ")
P("ตารางแสดงการแบ่งความรับผิดชอบงานตามรายการที่กำหนด งานใดที่ทำหลายคน หัวหน้าทีมระบุผู้รับผิดชอบหลัก")
TABLE([("สมาชิก", "รวมงาน (ชิ้น)")] + [(f"{i}. {nm}", "") for i, (_, nm) in enumerate(MEMBERS, 1)],
      (4.27, 2.0), "จำนวนงานของสมาชิกแต่ละคน", 2)
from openpyxl import load_workbook  # noqa: E402
ws = load_workbook(RUBRIC).active
items, started = [], False
for r in ws.iter_rows(values_only=True):
    if r[1] == "รายการ" and r[6] == "ผู้รับผิดชอบ": started = True; continue
    if started and isinstance(r[0], int): items.append(r)
    if started and r[1] == "คะแนนทีมงาน (พิมพ์)": break
rows = [("ที่", "รายการ", "วัตถุประสงค์", "ขนาด", "หน่วย", "การตรวจสอบ", "ผู้รับผิดชอบ")]
for r in items:
    rows.append((str(r[0]), str(r[1]).strip(), str(r[2] or "").strip(), f"{r[3]:g}" if r[3] is not None else "", str(r[4] or ""),
                 str(r[5] or "").strip(), ""))
TABLE(rows, (0.35, 1.05, 1.15, 0.45, 0.5, 2.0, 0.77), "หน้าที่รับผิดชอบ", 2, size=10)
H(3, "2.27.1 คะแนนการทำงานของทีม")
P("หัวหน้าทีมเป็นผู้ให้คะแนนการทำงานของสมาชิกแต่ละคน คะแนนเต็ม 5")
TABLE([("สมาชิก", "คะแนนที่ได้ (เต็ม 5)", "เหตุผลที่ได้คะแนน")] + [(f"{i}. {nm}", "", "") for i, (_, nm) in enumerate(MEMBERS, 1)],
      (2.4, 1.3, 2.57), "คะแนนการทำงานของทีม", 2)

# ------------------------------------------------------------------ เลขหน้า
s0 = doc.sections[0]; s0.different_first_page_header_footer = True
fp = s0.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; field(fp, "PAGE", "1")
for s in doc.sections[1:]: s.footer.is_linked_to_previous = True; s.header.is_linked_to_previous = True

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT); print("saved", OUT)
